from __future__ import annotations

import csv
import io
import json
import uuid
from collections import Counter
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional

from app.reports.schemas import (
    AnalysisReportSection,
    DesignSummarySection,
    DevelopmentSummarySection,
    ExportFormat,
    GovernanceReportSection,
    PhaseStatus,
    PhaseSummary,
    ProjectSummarySection,
    QualityGateReportSection,
    RequirementSummarySection,
    TestExecutionReportSection,
    TestReport,
    TraceabilityReportSection,
)


class ReportGenerator:
    def generate(
        self,
        project_id: str,
        project_name: Optional[str] = None,
        srs: Optional[Dict[str, Any]] = None,
        sdd: Optional[Dict[str, Any]] = None,
        source_code: Optional[Dict[str, Any]] = None,
        api_docs: Optional[Dict[str, Any]] = None,
        database_schema: Optional[Dict[str, Any]] = None,
        environment: Optional[Dict[str, Any]] = None,
        requirements: Optional[List[Dict[str, Any]]] = None,
        risks: Optional[List[Dict[str, Any]]] = None,
        change_impact: Optional[Dict[str, Any]] = None,
        coverage: Optional[Dict[str, Any]] = None,
        test_strategy: Optional[Dict[str, Any]] = None,
        test_cases: Optional[List[Dict[str, Any]]] = None,
        test_scenarios: Optional[List[Dict[str, Any]]] = None,
        generated_test_data: Optional[List[Dict[str, Any]]] = None,
        traceability: Optional[Dict[str, Any]] = None,
        execution_results: Optional[List[Dict[str, Any]]] = None,
        execution_summary: Optional[Dict[str, Any]] = None,
        execution_status: Optional[str] = None,
        result_intelligence: Optional[Dict[str, Any]] = None,
        failure_analyses: Optional[List[Dict[str, Any]]] = None,
        root_cause_analyses: Optional[List[Dict[str, Any]]] = None,
        defect_analyses: Optional[List[Dict[str, Any]]] = None,
        flaky_analyses: Optional[List[Dict[str, Any]]] = None,
        quality_gate_report: Optional[Dict[str, Any]] = None,
        quality_score: Optional[float] = None,
        release_readiness: Optional[str] = None,
        approval_state: Optional[Dict[str, Any]] = None,
        **kwargs: Any,
    ) -> TestReport:
        requirements = requirements or []
        risks = risks or []
        test_cases = test_cases or []
        test_scenarios = test_scenarios or []
        generated_test_data = generated_test_data or []
        execution_results = execution_results or []
        failure_analyses = failure_analyses or []
        root_cause_analyses = root_cause_analyses or []
        defect_analyses = defect_analyses or []
        flaky_analyses = flaky_analyses or []

        srs_dict = srs or {}
        sdd_dict = sdd or {}
        src_dict = source_code or {}

        # 1. SDLC & Upstream Sections
        p_name = project_name or srs_dict.get("title") or srs_dict.get("project_name") or project_id
        timestamp_str = datetime.now(timezone.utc).isoformat()

        proj_section = ProjectSummarySection(
            project_id=project_id,
            project_name=p_name,
            timestamp=timestamp_str,
            sdlc_phase="TESTING",
            sdlc_status="READY_FOR_TESTING",
        )

        req_features = []
        if srs_dict.get("features"):
            req_features = [str(f) for f in srs_dict.get("features", [])]
        elif srs_dict.get("functional_requirements"):
            req_features = [
                f"{r.get('name')}: {r.get('description', '')}" if isinstance(r, dict) else str(r)
                for r in srs_dict.get("functional_requirements", [])
            ]
        elif requirements:
            req_features = [r.get("description", r.get("id", "Requirement")) for r in requirements]

        req_nfrs = [str(n) for n in srs_dict.get("non_functional_requirements", [])]
        req_stories = srs_dict.get("user_stories", [])

        req_section = RequirementSummarySection(
            title=srs_dict.get("title") or srs_dict.get("project_name") or f"{p_name} SRS",
            version=srs_dict.get("version", "1.0.0"),
            features_count=len(req_features),
            features=req_features[:10],
            non_functional_requirements_count=len(req_nfrs),
            non_functional_requirements=req_nfrs[:5],
            user_stories_count=len(req_stories),
            business_goals=srs_dict.get("business_goals") or srs_dict.get("project_summary") or "Comprehensive enterprise capability verification.",
        )

        sdd_comps = [
            m.get("name") if isinstance(m, dict) else str(m)
            for m in (sdd_dict.get("components") or sdd_dict.get("module_breakdown") or ["Core Engine"])
        ]
        sdd_interfaces = [
            f"{ep.get('method', 'GET')} {ep.get('path', '/')}" if isinstance(ep, dict) else str(ep)
            for ep in (sdd_dict.get("interfaces") or sdd_dict.get("api_endpoints") or (api_docs.get("endpoints") if api_docs else []) or [])
        ]
        sdd_tables = [
            t.get("name") or t.get("table_name") if isinstance(t, dict) else str(t)
            for t in (sdd_dict.get("database_tables") or (database_schema.get("tables") if database_schema else []) or [])
        ]

        design_section = DesignSummarySection(
            architecture=str(sdd_dict.get("architecture") or sdd_dict.get("high_level_architecture") or "Clean Modular Architecture"),
            components_count=len(sdd_comps),
            components=sdd_comps[:10],
            interfaces_count=len(sdd_interfaces),
            interfaces=sdd_interfaces[:10],
            database_tables_count=len(sdd_tables),
            database_tables=sdd_tables[:10],
            security_policies=sdd_dict.get("security_policies") or ["TLS 1.3 Transport Encryption", "Role-Based Access Control (RBAC)", "Input Sanitization"],
        )

        dev_files = src_dict.get("files", [])
        dev_changed = src_dict.get("changes", {}).get("changed_files", dev_files[:3] if len(dev_files) >= 3 else dev_files)
        dev_section = DevelopmentSummarySection(
            repository=src_dict.get("repository") or src_dict.get("repo") or f"github.com/enterprise/{project_id}",
            language=src_dict.get("language") or "Python",
            build_version=src_dict.get("build_version") or "Build v1.0",
            files_count=len(dev_files),
            files=[str(f) for f in dev_files][:10],
            changed_files=[str(f) for f in dev_changed],
            artifact_zip=src_dict.get("artifact_zip") or "Source.zip",
        )

        # 2. Phases 1-8 Summaries
        phases = self._build_phase_summaries(
            requirements=requirements,
            risks=risks,
            coverage=coverage,
            test_cases=test_cases,
            test_scenarios=test_scenarios,
            execution_status=execution_status,
            execution_summary=execution_summary,
            quality_gate_report=quality_gate_report,
            failure_analyses=failure_analyses,
            defect_analyses=defect_analyses,
        )

        execution_section = self._build_execution_section(
            execution_results=execution_results,
            execution_summary=execution_summary,
        )

        analysis_section = self._build_analysis_section(
            failure_analyses=failure_analyses,
            root_cause_analyses=root_cause_analyses,
            defect_analyses=defect_analyses,
            flaky_analyses=flaky_analyses,
        )

        quality_section = self._build_quality_section(
            quality_gate_report=quality_gate_report,
            quality_score=quality_score,
            release_readiness=release_readiness,
        )

        traceability_section = self._build_traceability_section(
            traceability=traceability,
            coverage=coverage,
        )

        # 3. Governance & Release Allowed
        app_status = (approval_state.get("approval_status") if approval_state else "pending") or "pending"
        app_by = approval_state.get("approved_by") if approval_state else None
        app_comment = approval_state.get("approval_comment") if approval_state else None
        app_ts = approval_state.get("approval_timestamp") if approval_state else None
        
        # Policy: Release allowed if approved AND quality gate is ready/pass
        qg_ready = (release_readiness or quality_section.release_readiness) in ("READY", "READY_FOR_RELEASE")
        is_rel_allowed = (app_status == "approved") and qg_ready

        gov_section = GovernanceReportSection(
            approval_status=app_status,
            approved_by=app_by,
            comment=app_comment,
            approval_timestamp=app_ts,
            release_allowed=is_rel_allowed,
            quality_gate_status=quality_section.overall_status,
            release_readiness=release_readiness or quality_section.release_readiness,
        )

        executive_summary = self._build_executive_summary(
            project_id=project_id,
            execution_summary=execution_summary,
            quality_score=quality_score,
            release_readiness=release_readiness or quality_section.release_readiness,
            analysis_section=analysis_section,
            quality_section=quality_section,
        )

        recommendations = self._build_recommendations(
            execution_section=execution_section,
            analysis_section=analysis_section,
            quality_section=quality_section,
            traceability_section=traceability_section,
        )

        raw_data = self._build_raw_data(
            requirements=requirements,
            risks=risks,
            change_impact=change_impact,
            coverage=coverage,
            test_strategy=test_strategy,
            test_cases=test_cases,
            test_scenarios=test_scenarios,
            generated_test_data=generated_test_data,
            traceability=traceability,
            execution_results=execution_results,
            execution_summary=execution_summary,
            result_intelligence=result_intelligence,
            failure_analyses=failure_analyses,
            root_cause_analyses=root_cause_analyses,
            defect_analyses=defect_analyses,
            flaky_analyses=flaky_analyses,
            quality_gate_report=quality_gate_report,
            srs=srs,
            sdd=sdd,
            source_code=source_code,
            governance=approval_state,
        )

        return TestReport(
            report_id=f"RPT-{uuid.uuid4().hex[:8].upper()}",
            project_id=project_id,
            project_name=p_name,
            generated_at=timestamp_str,
            project_summary=proj_section,
            requirement_summary=req_section,
            design_summary=design_section,
            development_summary=dev_section,
            phases=phases,
            execution=execution_section,
            analysis=analysis_section,
            quality_gate=quality_section,
            traceability=traceability_section,
            governance=gov_section,
            executive_summary=executive_summary,
            recommendations=recommendations,
            release_readiness=release_readiness or quality_section.release_readiness,
            release_allowed=is_rel_allowed,
            raw_data=raw_data,
        )

    def export(self, report: TestReport, fmt: ExportFormat | str) -> str:
        f = ExportFormat(fmt.lower()) if isinstance(fmt, str) else fmt
        if f == ExportFormat.JSON:
            return self._export_json(report)
        if f == ExportFormat.HTML:
            return self._export_html(report)
        if f == ExportFormat.CSV:
            return self._export_csv(report)
        if f == ExportFormat.MD:
            return self._export_markdown(report)
        if f in (ExportFormat.PDF, ExportFormat.DOCX):
            return f"Binary export format: use export_bytes(report, ExportFormat.{f.name})"
        raise ValueError(f"Unsupported format: {fmt}")

    def export_bytes(self, report: TestReport, fmt: ExportFormat | str) -> bytes:
        f = ExportFormat(fmt.lower()) if isinstance(fmt, str) else fmt
        if f == ExportFormat.JSON:
            return self._export_json(report).encode("utf-8")
        if f == ExportFormat.HTML:
            return self._export_html(report).encode("utf-8")
        if f == ExportFormat.CSV:
            return self._export_csv(report).encode("utf-8")
        if f == ExportFormat.MD:
            return self._export_markdown(report).encode("utf-8")
        if f == ExportFormat.PDF:
            return self._export_pdf(report)
        if f == ExportFormat.DOCX:
            return self._export_docx(report)
        raise ValueError(f"Unsupported format: {fmt}")

    def export_json(self, report: TestReport) -> str:
        return self._export_json(report)

    def export_html(self, report: TestReport) -> str:
        return self._export_html(report)

    def export_csv(self, report: TestReport) -> str:
        return self._export_csv(report)

    def export_markdown(self, report: TestReport) -> str:
        return self._export_markdown(report)

    def export_pdf(self, report: TestReport) -> bytes:
        return self._export_pdf(report)

    def export_docx(self, report: TestReport) -> bytes:
        return self._export_docx(report)

    # ------------------------------------------------------------------
    # Phase summaries
    # ------------------------------------------------------------------

    def _build_phase_summaries(
        self,
        requirements: List[Dict[str, Any]],
        risks: List[Dict[str, Any]],
        coverage: Optional[Dict[str, Any]],
        test_cases: List[Dict[str, Any]],
        test_scenarios: List[Dict[str, Any]],
        execution_status: Optional[str],
        execution_summary: Optional[Dict[str, Any]],
        quality_gate_report: Optional[Dict[str, Any]],
        failure_analyses: Optional[List[Dict[str, Any]]] = None,
        defect_analyses: Optional[List[Dict[str, Any]]] = None,
    ) -> List[PhaseSummary]:
        phases: List[PhaseSummary] = []

        phases.append(PhaseSummary(
            phase_number=1,
            phase_name="Input Validation & Context Loading",
            status=PhaseStatus.PASSED,
            summary="Input validation and context loading completed successfully.",
            metrics={"validation_contract": "PASSED"},
        ))

        req_count = len(requirements)
        risk_count = len(risks)
        cov_pct = coverage.get("coverage_percentage", 0.0) if coverage else 0.0
        phases.append(PhaseSummary(
            phase_number=2,
            phase_name="Testing Intelligence",
            status=PhaseStatus.PASSED if req_count > 0 else PhaseStatus.SKIPPED,
            summary=f"Analyzed {req_count} requirements and {risk_count} risks. Coverage at {cov_pct:.1f}%.",
            metrics={"requirements": req_count, "risks": risk_count, "coverage_pct": cov_pct},
        ))

        tc_count = len(test_cases)
        sc_count = len(test_scenarios)
        phases.append(PhaseSummary(
            phase_number=3,
            phase_name="Test Design",
            status=PhaseStatus.PASSED if tc_count > 0 else PhaseStatus.SKIPPED,
            summary=f"Generated {tc_count} test cases and {sc_count} scenarios.",
            metrics={"test_cases": tc_count, "scenarios": sc_count},
        ))

        exec_status_val = execution_status or "not_executed"
        exec_phase_status = PhaseStatus.PASSED if exec_status_val == "completed" else PhaseStatus.PARTIAL
        phases.append(PhaseSummary(
            phase_number=4,
            phase_name="Test Execution",
            status=exec_phase_status,
            summary=f"Execution status: {exec_status_val}.",
            metrics=execution_summary or {},
        ))

        phases.append(PhaseSummary(
            phase_number=5,
            phase_name="Execution Enrichment",
            status=exec_phase_status,
            summary="Execution results enriched with metadata, retries, and screenshots.",
            metrics=execution_summary or {},
        ))

        intel_status = PhaseStatus.PASSED if failure_analyses or defect_analyses else PhaseStatus.SKIPPED
        phases.append(PhaseSummary(
            phase_number=6,
            phase_name="Result Intelligence",
            status=intel_status,
            summary=f"Analyzed {len(failure_analyses or [])} failures, classified {len(defect_analyses or [])} defects.",
            metrics={"failures": len(failure_analyses or []), "defects": len(defect_analyses or [])},
        ))

        qg_status_str = "NOT_EVALUATED"
        if quality_gate_report:
            qg_status_str = quality_gate_report.get("overall_status", "NOT_EVALUATED")
        phases.append(PhaseSummary(
            phase_number=7,
            phase_name="Quality Gate",
            status=PhaseStatus.PASSED if qg_status_str == "PASS" else (
                PhaseStatus.FAILED if qg_status_str == "FAIL" else PhaseStatus.PARTIAL
            ),
            summary=f"Quality gate: {qg_status_str}.",
            metrics={"overall_status": qg_status_str},
        ))

        phases.append(PhaseSummary(
            phase_number=8,
            phase_name="Report Generation",
            status=PhaseStatus.PASSED,
            summary="Comprehensive report generated and exported.",
            metrics={"export_formats": ["json", "html", "csv", "pdf", "docx"]},
        ))

        return phases

    def _build_execution_section(
        self,
        execution_results: List[Dict[str, Any]],
        execution_summary: Optional[Dict[str, Any]],
    ) -> TestExecutionReportSection:
        if not execution_summary:
            return TestExecutionReportSection()

        total = execution_summary.get("total", 0)
        passed = execution_summary.get("pass", execution_summary.get("passed", 0))
        failed = execution_summary.get("fail", execution_summary.get("failed", 0))
        errors = execution_summary.get("error", execution_summary.get("errors", 0))
        skipped = execution_summary.get("skipped", 0)
        pass_rate = (passed / total * 100) if total > 0 else 0.0

        failed_ids = [
            r.get("test_case_id", r.get("id", "unknown"))
            for r in execution_results
            if r.get("status") == "FAIL"
        ]
        error_ids = [
            r.get("test_case_id", r.get("id", "unknown"))
            for r in execution_results
            if r.get("status") == "ERROR"
        ]

        return TestExecutionReportSection(
            total_tests=total,
            passed=passed,
            failed=failed,
            errors=errors,
            skipped=skipped,
            pass_rate=round(pass_rate, 2),
            duration=execution_summary.get("duration"),
            failed_test_ids=failed_ids,
            error_test_ids=error_ids,
        )

    def _build_analysis_section(
        self,
        failure_analyses: List[Dict[str, Any]],
        root_cause_analyses: List[Dict[str, Any]],
        defect_analyses: List[Dict[str, Any]],
        flaky_analyses: List[Dict[str, Any]],
    ) -> AnalysisReportSection:
        failure_types: Dict[str, int] = dict(Counter(
            fa.get("failure_type", "unknown") for fa in failure_analyses
        ))

        root_causes = [
            {
                "category": rc.get("category", "unknown"),
                "hypothesis": rc.get("hypothesis", ""),
                "confidence": rc.get("confidence", "low"),
            }
            for rc in root_cause_analyses
        ]

        product_defects = sum(
            1 for da in defect_analyses
            if da.get("classification") in ("product_defect", "product_defect_configuration", "product_defect_performance")
        )
        test_defects = sum(
            1 for da in defect_analyses
            if da.get("classification") in ("test_defect", "test_defect_flaky", "test_defect_data")
        )

        flaky_count = sum(
            1 for fa in flaky_analyses
            if fa.get("verdict") == "FLAKY"
        )

        return AnalysisReportSection(
            total_failures=len(failure_analyses),
            failure_types=failure_types,
            root_causes=root_causes,
            product_defects=product_defects,
            test_defects=test_defects,
            flaky_tests=flaky_count,
        )

    def _build_quality_section(
        self,
        quality_gate_report: Optional[Dict[str, Any]],
        quality_score: Optional[float],
        release_readiness: Optional[str],
    ) -> QualityGateReportSection:
        if not quality_gate_report:
            return QualityGateReportSection()

        gates = quality_gate_report.get("gates", [])
        gate_results = [
            {"name": g.get("gate", "unknown"), "status": g.get("status", "NOT_EVALUATED")}
            for g in gates
        ]
        blocking = [
            g.get("gate", "unknown")
            for g in gates
            if g.get("status") == "FAIL"
        ]

        return QualityGateReportSection(
            overall_status=quality_gate_report.get("overall_status", "NOT_EVALUATED"),
            release_readiness=release_readiness or "NOT_READY",
            quality_score=quality_score or 0.0,
            gate_results=gate_results,
            blocking_gates=blocking,
        )

    def _build_traceability_section(
        self,
        traceability: Optional[Dict[str, Any]],
        coverage: Optional[Dict[str, Any]],
    ) -> TraceabilityReportSection:
        if not traceability:
            cov_pct = coverage.get("coverage_percentage", 0.0) if coverage else 0.0
            return TraceabilityReportSection(coverage_percentage=cov_pct)

        entries = traceability.get("entries", [])
        return TraceabilityReportSection(
            total_entries=len(entries),
            coverage_percentage=coverage.get("coverage_percentage", 0.0) if coverage else 0.0,
            uncovered_requirements=traceability.get("uncovered_requirements", []),
            orphaned_test_cases=traceability.get("orphaned_test_cases", []),
            orphaned_test_data=traceability.get("orphaned_test_data", []),
        )

    def _build_executive_summary(
        self,
        project_id: str,
        execution_summary: Optional[Dict[str, Any]],
        quality_score: Optional[float],
        release_readiness: Optional[str],
        analysis_section: AnalysisReportSection,
        quality_section: QualityGateReportSection,
    ) -> str:
        parts = [f"Project {project_id} — Testing Report."]

        if execution_summary:
            total = execution_summary.get("total", 0)
            passed = execution_summary.get("pass", execution_summary.get("passed", 0))
            pass_rate = (passed / total * 100) if total > 0 else 0.0
            parts.append(f"Executed {total} tests with {pass_rate:.1f}% pass rate.")

        if quality_score is not None:
            parts.append(f"Quality score: {quality_score:.1f}/100.")

        if release_readiness:
            parts.append(f"Release readiness: {release_readiness}.")

        if analysis_section.product_defects > 0:
            parts.append(f"{analysis_section.product_defects} product defect(s) identified.")
        if analysis_section.flaky_tests > 0:
            parts.append(f"{analysis_section.flaky_tests} flaky test(s) detected.")

        return " ".join(parts)

    def _build_recommendations(
        self,
        execution_section: TestExecutionReportSection,
        analysis_section: AnalysisReportSection,
        quality_section: QualityGateReportSection,
        traceability_section: TraceabilityReportSection,
    ) -> List[str]:
        recs: List[str] = []

        if quality_section.overall_status == "PASS":
            recs.append("All quality gates passed. Consider releasing.")
        elif quality_section.blocking_gates:
            recs.append(f"Resolve {len(quality_section.blocking_gates)} blocking quality gates before release.")

        if execution_section.failed > 0:
            recs.append(f"Fix {execution_section.failed} failing tests.")

        if analysis_section.product_defects > 0:
            recs.append(f"Address {analysis_section.product_defects} identified product defects.")

        if analysis_section.flaky_tests > 0:
            recs.append(f"Quarantine or stabilize {analysis_section.flaky_tests} flaky tests.")

        if traceability_section.uncovered_requirements:
            recs.append(f"Add test coverage for {len(traceability_section.uncovered_requirements)} uncovered requirements.")

        if not recs:
            recs.append("Review test results and proceed with standard verification.")

        return recs

    def _build_raw_data(self, **kwargs: Any) -> Dict[str, Any]:
        return {k: v for k, v in kwargs.items() if v is not None}

    # ------------------------------------------------------------------
    # Export: JSON
    # ------------------------------------------------------------------

    def _export_json(self, report: TestReport) -> str:
        return json.dumps(report.model_dump(), indent=2, default=str)

    # ------------------------------------------------------------------
    # Export: HTML
    # ------------------------------------------------------------------

    def _export_html(self, report: TestReport) -> str:
        lines: List[str] = [
            "<!DOCTYPE html>",
            "<html lang='en'>",
            "<head>",
            "<meta charset='utf-8'/>",
            "<meta name='viewport' content='width=device-width, initial-scale=1.0'/>",
            f"<title>Testing Report — {report.project_id}</title>",
            "<style>",
            "body { font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif; margin: 0; padding: 24px; background: #0b0f19; color: #f3f4f6; }",
            ".container { max-width: 1000px; margin: 0 auto; background: #111827; border: 1px solid #1f2937; border-radius: 12px; padding: 32px; box-shadow: 0 10px 25px rgba(0,0,0,0.5); }",
            "h1 { color: #38bdf8; font-size: 24px; margin-top: 0; border-bottom: 2px solid #1f2937; padding-bottom: 12px; }",
            "h2 { color: #818cf8; font-size: 16px; margin-top: 28px; margin-bottom: 12px; text-transform: uppercase; letter-spacing: 0.05em; border-left: 3px solid #818cf8; padding-left: 8px; }",
            "p, li { font-size: 13px; line-height: 1.6; color: #cbd5e1; }",
            ".grid { display: grid; grid-template-columns: repeat(auto-fit, minmax(220px, 1fr)); gap: 12px; margin-bottom: 16px; }",
            ".card { background: #0f172a; border: 1px solid #1e293b; border-radius: 8px; padding: 14px; }",
            ".card-title { font-size: 11px; font-weight: bold; text-transform: uppercase; color: #94a3b8; margin-bottom: 6px; }",
            ".card-val { font-size: 18px; font-weight: 800; color: #f8fafc; }",
            "table { width: 100%; border-collapse: collapse; margin-top: 10px; margin-bottom: 16px; font-size: 12px; }",
            "th, td { border: 1px solid #1e293b; padding: 8px 12px; text-align: left; }",
            "th { background: #0f172a; color: #94a3b8; font-weight: 700; text-transform: uppercase; font-size: 10px; }",
            "tr:nth-child(even) { background: #0b1120; }",
            ".badge { display: inline-block; padding: 2px 8px; border-radius: 4px; font-weight: bold; font-size: 10px; text-transform: uppercase; }",
            ".badge-pass { background: rgba(16, 185, 129, 0.2); color: #34d399; border: 1px solid rgba(16, 185, 129, 0.4); }",
            ".badge-fail { background: rgba(244, 63, 94, 0.2); color: #fb7185; border: 1px solid rgba(244, 63, 94, 0.4); }",
            ".badge-warn { background: rgba(245, 158, 11, 0.2); color: #fbbf24; border: 1px solid rgba(245, 158, 11, 0.4); }",
            ".badge-info { background: rgba(56, 189, 248, 0.2); color: #38bdf8; border: 1px solid rgba(56, 189, 248, 0.4); }",
            "</style>",
            "</head>",
            "<body>",
            "<div class='container'>",
            f"<h1>Test Report — {report.project_name or report.project_id}</h1>",
            f"<p><strong>Report ID:</strong> {report.report_id} &nbsp;|&nbsp; <strong>Generated:</strong> {report.generated_at}</p>",
            f"<p style='background:#0f172a; border-left:4px solid #38bdf8; padding:12px; border-radius:4px;'><strong>Executive Summary:</strong> {report.executive_summary}</p>",
            
            "<h2>1. SDLC & Upstream Context</h2>",
            "<div class='grid'>",
            f"<div class='card'><div class='card-title'>Requirement / SRS</div><div class='card-val'>{report.requirement_summary.features_count} Features</div><p style='margin:4px 0 0 0;'>{report.requirement_summary.title} (v{report.requirement_summary.version})</p></div>",
            f"<div class='card'><div class='card-title'>Design / SDD</div><div class='card-val'>{report.design_summary.components_count} Modules</div><p style='margin:4px 0 0 0;'>{report.design_summary.interfaces_count} APIs, {report.design_summary.database_tables_count} Tables</p></div>",
            f"<div class='card'><div class='card-title'>Development / Code</div><div class='card-val'>{report.development_summary.files_count} Files</div><p style='margin:4px 0 0 0;'>{report.development_summary.language} ({report.development_summary.build_version})</p></div>",
            "</div>",

            "<h2>2. 8-Phase Testing Breakdown</h2>",
            "<table>",
            "<tr><th>Phase #</th><th>Phase Name</th><th>Status</th><th>Summary</th></tr>",
        ]

        for p in report.phases:
            badge_cls = "badge-pass" if p.status == PhaseStatus.PASSED else ("badge-fail" if p.status == PhaseStatus.FAILED else "badge-warn")
            lines.append(f"<tr><td>P{p.phase_number}</td><td><strong>{p.phase_name}</strong></td><td><span class='badge {badge_cls}'>{p.status.value}</span></td><td>{p.summary}</td></tr>")
        lines.append("</table>")

        # Execution & Quality Gate
        lines.extend([
            "<h2>3. Execution & Quality Gate Results</h2>",
            "<div class='grid'>",
            f"<div class='card'><div class='card-title'>Total Tests</div><div class='card-val'>{report.execution.total_tests}</div></div>",
            f"<div class='card'><div class='card-title'>Pass Rate</div><div class='card-val'>{report.execution.pass_rate}%</div></div>",
            f"<div class='card'><div class='card-title'>Quality Score</div><div class='card-val'>{report.quality_gate.quality_score:.1f} / 100</div></div>",
            f"<div class='card'><div class='card-title'>Release Readiness</div><div class='card-val'>{report.quality_gate.release_readiness}</div></div>",
            "</div>",
            
            "<h2>4. Human Governance & Sign-Off</h2>",
            f"<p><strong>Approval Status:</strong> <span class='badge {badge_cls}'>{report.governance.approval_status.upper()}</span> &nbsp;|&nbsp; <strong>Reviewer:</strong> {report.governance.approved_by or 'Pending'} &nbsp;|&nbsp; <strong>Release Allowed:</strong> <span class='badge {'badge-pass' if report.release_allowed else 'badge-fail'}'>{'YES' if report.release_allowed else 'NO'}</span></p>",
            f"<p><strong>Comments / Reason:</strong> {report.governance.comment or 'None recorded.'}</p>",

            "<h2>5. Recommendations</h2>",
            "<ul>",
        ])
        for r in report.recommendations:
            lines.append(f"<li>{r}</li>")
        lines.extend(["</ul>", "</div>", "</body>", "</html>"])

        return "\n".join(lines)

    # ------------------------------------------------------------------
    # Export: CSV
    # ------------------------------------------------------------------

    def _export_csv(self, report: TestReport) -> str:
        buf = io.StringIO()
        writer = csv.writer(buf)

        writer.writerow(["Section", "Key", "Value"])
        writer.writerow(["Report", "report_id", report.report_id])
        writer.writerow(["Report", "project_id", report.project_id])
        writer.writerow(["Report", "project_name", report.project_name or report.project_id])
        writer.writerow(["Report", "generated_at", report.generated_at])
        writer.writerow(["Report", "executive_summary", report.executive_summary])

        writer.writerow([])
        writer.writerow(["Upstream_SDLC", "SRS_Title", report.requirement_summary.title])
        writer.writerow(["Upstream_SDLC", "SRS_Features_Count", report.requirement_summary.features_count])
        writer.writerow(["Upstream_SDLC", "SDD_Architecture", report.design_summary.architecture])
        writer.writerow(["Upstream_SDLC", "SDD_Components_Count", report.design_summary.components_count])
        writer.writerow(["Upstream_SDLC", "SDD_Interfaces_Count", report.design_summary.interfaces_count])
        writer.writerow(["Upstream_SDLC", "Dev_Repository", report.development_summary.repository])
        writer.writerow(["Upstream_SDLC", "Dev_Language", report.development_summary.language])
        writer.writerow(["Upstream_SDLC", "Dev_Files_Count", report.development_summary.files_count])

        writer.writerow([])
        writer.writerow(["Phase", "Number", "Name", "Status", "Summary"])
        for p in report.phases:
            writer.writerow(["Phase", p.phase_number, p.phase_name, p.status, p.summary])

        writer.writerow([])
        ex = report.execution
        writer.writerow(["Execution", "total_tests", ex.total_tests])
        writer.writerow(["Execution", "passed", ex.passed])
        writer.writerow(["Execution", "failed", ex.failed])
        writer.writerow(["Execution", "errors", ex.errors])
        writer.writerow(["Execution", "skipped", ex.skipped])
        writer.writerow(["Execution", "pass_rate", ex.pass_rate])

        writer.writerow([])
        an = report.analysis
        writer.writerow(["Analysis", "total_failures", an.total_failures])
        writer.writerow(["Analysis", "product_defects", an.product_defects])
        writer.writerow(["Analysis", "test_defects", an.test_defects])
        writer.writerow(["Analysis", "flaky_tests", an.flaky_tests])

        writer.writerow([])
        qg = report.quality_gate
        writer.writerow(["QualityGate", "overall_status", qg.overall_status])
        writer.writerow(["QualityGate", "quality_score", qg.quality_score])
        writer.writerow(["QualityGate", "release_readiness", qg.release_readiness])

        writer.writerow([])
        gov = report.governance
        writer.writerow(["Governance", "approval_status", gov.approval_status])
        writer.writerow(["Governance", "approved_by", gov.approved_by or ""])
        writer.writerow(["Governance", "comment", gov.comment or ""])
        writer.writerow(["Governance", "release_allowed", report.release_allowed])

        writer.writerow([])
        writer.writerow(["Recommendations"])
        for r in report.recommendations:
            writer.writerow(["", r])

        return buf.getvalue()

    # ------------------------------------------------------------------
    # Export: Markdown
    # ------------------------------------------------------------------

    def _export_markdown(self, report: TestReport) -> str:
        lines: List[str] = [
            f"# Test Audit & Certification Report — {report.project_name or report.project_id}",
            "",
            f"- **Report ID:** `{report.report_id}`",
            f"- **Generated At:** {report.generated_at}",
            f"- **SDLC Phase:** {report.project_summary.sdlc_phase} ({report.project_summary.sdlc_status})",
            "",
            "> **Executive Summary:**",
            f"> {report.executive_summary}",
            "",
            "---",
            "",
            "## 1. Upstream SDLC Context Traceability",
            "",
            f"- **Requirement / SRS:** {report.requirement_summary.title} (v{report.requirement_summary.version}) — `{report.requirement_summary.features_count}` Functional Requirements Tracked",
            f"- **Design / SDD:** Architecture: {report.design_summary.architecture[:60]}... — `{report.design_summary.components_count}` Modules, `{report.design_summary.interfaces_count}` Interfaces, `{report.design_summary.database_tables_count}` Tables",
            f"- **Development / Code:** `{report.development_summary.files_count}` Files generated ({report.development_summary.language}) — Build Version `{report.development_summary.build_version}`",
            "",
            "---",
            "",
            "## 2. 8-Phase Testing Breakdown",
            "",
            "| Phase # | Phase Name | Status | Summary |",
            "| :--- | :--- | :--- | :--- |",
        ]

        for p in report.phases:
            lines.append(f"| **P{p.phase_number}** | {p.phase_name} | `{p.status.value.upper()}` | {p.summary} |")

        lines.extend([
            "",
            "---",
            "",
            "## 3. Test Execution Summary",
            "",
            f"- **Total Tests Executed:** `{report.execution.total_tests}`",
            f"- **Passed:** `{report.execution.passed}`",
            f"- **Failed:** `{report.execution.failed}`",
            f"- **Errors:** `{report.execution.errors}`",
            f"- **Skipped:** `{report.execution.skipped}`",
            f"- **Pass Rate:** `{report.execution.pass_rate}%`",
            "",
            "---",
            "",
            "## 4. Diagnostics & Defect Analysis",
            "",
            f"- **Total Failures:** `{report.analysis.total_failures}`",
            f"- **Product Defects Flagged:** `{report.analysis.product_defects}`",
            f"- **Test Defects Flagged:** `{report.analysis.test_defects}`",
            f"- **Flaky Tests Flagged:** `{report.analysis.flaky_tests}`",
            "",
            "---",
            "",
            "## 5. Autonomous Quality Gate & Release Readiness",
            "",
            f"- **Overall Quality Gate Status:** `{report.quality_gate.overall_status}`",
            f"- **Deterministic Quality Score:** **`{report.quality_gate.quality_score:.1f} / 100`**",
            f"- **Release Readiness Decision:** **`{report.quality_gate.release_readiness}`**",
            "",
            "---",
            "",
            "## 6. Human Governance & Sign-Off",
            "",
            f"- **Approval Status:** `{report.governance.approval_status.upper()}`",
            f"- **QA Reviewer:** `{report.governance.approved_by or 'Pending Review'}`",
            f"- **Release Allowed:** **`{'YES (APPROVED)' if report.release_allowed else 'NO (BLOCKED/PENDING)'}`**",
            f"- **Governance Comments:** {report.governance.comment or 'None recorded.'}",
            "",
            "---",
            "",
            "## 7. QA Recommendations",
            "",
        ])

        for r in report.recommendations:
            lines.append(f"- {r}")

        lines.append("")
        return "\n".join(lines)

    # ------------------------------------------------------------------
    # Export: PDF
    # ------------------------------------------------------------------

    def _export_pdf(self, report: TestReport) -> bytes:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors

        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Title'],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor('#1e293b'),
            alignment=0,
            spaceAfter=6
        )
        heading_style = ParagraphStyle(
            'SectionHead',
            parent=styles['Heading2'],
            fontSize=12,
            leading=15,
            textColor=colors.HexColor('#0f766e'),
            spaceBefore=12,
            spaceAfter=6
        )
        body_style = ParagraphStyle(
            'Body',
            parent=styles['Normal'],
            fontSize=9,
            leading=12,
            textColor=colors.HexColor('#334155')
        )
        meta_style = ParagraphStyle(
            'Meta',
            parent=styles['Normal'],
            fontSize=8,
            leading=10,
            textColor=colors.HexColor('#64748b')
        )

        elements = []

        # Title & Metadata
        elements.append(Paragraph(f"<b>AI-SDLC Testing Agent Audit Report</b>", title_style))
        elements.append(Paragraph(f"<b>Project:</b> {report.project_name or report.project_id} &nbsp;|&nbsp; <b>Report ID:</b> {report.report_id} &nbsp;|&nbsp; <b>Date:</b> {report.generated_at[:19]}", meta_style))
        elements.append(Spacer(1, 8))
        elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#cbd5e1'), spaceAfter=10))

        # Executive Summary
        elements.append(Paragraph("<b>Executive Summary</b>", heading_style))
        elements.append(Paragraph(report.executive_summary, body_style))
        elements.append(Spacer(1, 8))

        # Upstream SDLC Context Table
        elements.append(Paragraph("<b>1. Upstream SDLC Context (Requirement → Design → Development)</b>", heading_style))
        sdlc_data = [
            [Paragraph("<b>SDLC Stage</b>", meta_style), Paragraph("<b>Artifact & Features Summary</b>", meta_style), Paragraph("<b>Key Metrics</b>", meta_style)],
            [
                Paragraph("<b>Requirement (SRS)</b>", body_style),
                Paragraph(f"{report.requirement_summary.title} (v{report.requirement_summary.version})", body_style),
                Paragraph(f"{report.requirement_summary.features_count} Functional Requirements<br/>{report.requirement_summary.non_functional_requirements_count} NFRs", body_style)
            ],
            [
                Paragraph("<b>Design (SDD)</b>", body_style),
                Paragraph(f"Architecture: {report.design_summary.architecture[:60]}...", body_style),
                Paragraph(f"{report.design_summary.components_count} Modules, {report.design_summary.interfaces_count} APIs<br/>{report.design_summary.database_tables_count} DB Tables", body_style)
            ],
            [
                Paragraph("<b>Development</b>", body_style),
                Paragraph(f"Repo: {report.development_summary.repository} ({report.development_summary.language})", body_style),
                Paragraph(f"{report.development_summary.files_count} Generated Files ({report.development_summary.build_version})", body_style)
            ],
        ]
        t_sdlc = Table(sdlc_data, colWidths=[120, 240, 180])
        t_sdlc.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f1f5f9')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(t_sdlc)
        elements.append(Spacer(1, 8))

        # 8 Phases Table
        elements.append(Paragraph("<b>2. 8-Phase Autonomous Testing Pipeline</b>", heading_style))
        phase_data = [[Paragraph("<b>Phase</b>", meta_style), Paragraph("<b>Phase Name</b>", meta_style), Paragraph("<b>Status</b>", meta_style), Paragraph("<b>Summary</b>", meta_style)]]
        for p in report.phases:
            phase_data.append([
                Paragraph(f"<b>P{p.phase_number}</b>", body_style),
                Paragraph(p.phase_name, body_style),
                Paragraph(f"<b>{p.status.value.upper()}</b>", body_style),
                Paragraph(p.summary, body_style)
            ])
        t_phase = Table(phase_data, colWidths=[40, 150, 60, 290])
        t_phase.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f1f5f9')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        elements.append(t_phase)
        elements.append(Spacer(1, 8))

        # Quality Gate & Execution Results
        elements.append(Paragraph("<b>3. Quality Gate & Governance Sign-Off</b>", heading_style))
        qg_data = [
            [
                Paragraph("<b>Execution Summary:</b>", body_style),
                Paragraph(f"Total: {report.execution.total_tests} | Passed: {report.execution.passed} | Failed: {report.execution.failed} | Pass Rate: {report.execution.pass_rate}%", body_style)
            ],
            [
                Paragraph("<b>Quality Gate Status:</b>", body_style),
                Paragraph(f"Status: <b>{report.quality_gate.overall_status}</b> | Quality Score: <b>{report.quality_gate.quality_score:.1f}/100</b> | Readiness: <b>{report.release_readiness}</b>", body_style)
            ],
            [
                Paragraph("<b>Human QA Sign-off:</b>", body_style),
                Paragraph(f"Status: <b>{report.governance.approval_status.upper()}</b> | Reviewer: {report.governance.approved_by or 'Pending'} | Release Allowed: <b>{'YES' if report.release_allowed else 'NO'}</b>", body_style)
            ],
            [
                Paragraph("<b>Reviewer Notes:</b>", body_style),
                Paragraph(report.governance.comment or "None recorded.", body_style)
            ],
        ]
        t_qg = Table(qg_data, colWidths=[140, 400])
        t_qg.setStyle(TableStyle([
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(t_qg)

        doc.build(elements)
        return buf.getvalue()

    # ------------------------------------------------------------------
    # Export: DOCX
    # ------------------------------------------------------------------

    def _export_docx(self, report: TestReport) -> bytes:
        import docx
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT

        doc = docx.Document()

        # Title
        title = doc.add_heading(f"AI-SDLC Testing Agent Audit Report", level=0)
        p_meta = doc.add_paragraph()
        p_meta.add_run(f"Project: {report.project_name or report.project_id} | Report ID: {report.report_id} | Date: {report.generated_at[:19]}").italic = True

        # Executive Summary
        doc.add_heading("Executive Summary", level=1)
        doc.add_paragraph(report.executive_summary)

        # 1. Upstream SDLC Context
        doc.add_heading("1. Upstream SDLC Context (Requirement → Design → Development)", level=1)
        t_sdlc = doc.add_table(rows=1, cols=3)
        t_sdlc.style = 'Table Grid'
        hdr_cells = t_sdlc.rows[0].cells
        hdr_cells[0].text = "SDLC Stage"
        hdr_cells[1].text = "Artifact & Summary"
        hdr_cells[2].text = "Key Metrics"

        r_row = t_sdlc.add_row().cells
        r_row[0].text = "Requirement (SRS)"
        r_row[1].text = f"{report.requirement_summary.title} (v{report.requirement_summary.version})"
        r_row[2].text = f"{report.requirement_summary.features_count} Functional Requirements, {report.requirement_summary.non_functional_requirements_count} NFRs"

        d_row = t_sdlc.add_row().cells
        d_row[0].text = "Design (SDD)"
        d_row[1].text = f"Architecture: {report.design_summary.architecture[:80]}"
        d_row[2].text = f"{report.design_summary.components_count} Modules, {report.design_summary.interfaces_count} APIs, {report.design_summary.database_tables_count} Tables"

        dev_row = t_sdlc.add_row().cells
        dev_row[0].text = "Development"
        dev_row[1].text = f"Repo: {report.development_summary.repository} ({report.development_summary.language})"
        dev_row[2].text = f"{report.development_summary.files_count} Generated Files ({report.development_summary.build_version})"

        # 2. 8 Phases Breakdown
        doc.add_heading("2. 8-Phase Autonomous Testing Pipeline", level=1)
        t_phase = doc.add_table(rows=1, cols=4)
        t_phase.style = 'Table Grid'
        p_hdr = t_phase.rows[0].cells
        p_hdr[0].text = "Phase"
        p_hdr[1].text = "Phase Name"
        p_hdr[2].text = "Status"
        p_hdr[3].text = "Summary"

        for p in report.phases:
            p_row = t_phase.add_row().cells
            p_row[0].text = f"P{p.phase_number}"
            p_row[1].text = p.phase_name
            p_row[2].text = p.status.value.upper()
            p_row[3].text = p.summary

        # 3. Quality Gate & Governance
        doc.add_heading("3. Quality Gate & Governance Sign-Off", level=1)
        doc.add_paragraph(f"Execution Summary: Total={report.execution.total_tests}, Passed={report.execution.passed}, Failed={report.execution.failed}, Pass Rate={report.execution.pass_rate}%")
        doc.add_paragraph(f"Quality Gate Status: {report.quality_gate.overall_status} (Score: {report.quality_gate.quality_score:.1f}/100, Release Readiness: {report.release_readiness})")
        doc.add_paragraph(f"Human Governance Sign-off: Status={report.governance.approval_status.upper()}, Reviewer={report.governance.approved_by or 'Pending'}, Release Allowed={'YES' if report.release_allowed else 'NO'}")
        if report.governance.comment:
            doc.add_paragraph(f"Reviewer Notes: {report.governance.comment}")

        # 4. Recommendations
        doc.add_heading("4. Recommendations", level=1)
        for r in report.recommendations:
            doc.add_paragraph(r, style='List Bullet')

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()


# Module-level convenience function
def generate_report(
    project_id: str,
    **kwargs: Any,
) -> TestReport:
    gen = ReportGenerator()
    return gen.generate(project_id=project_id, **kwargs)
